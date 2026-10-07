"""Document edits in the worker: formatting, references, and refusals that keep content."""

import io
import json
from collections.abc import Callable
from typing import Any
from uuid import uuid4

import pytest
from pydantic import TypeAdapter

from services.documents.operations.presentation import PresentationOperation
from services.documents.operations.word import WordOperation
from services.documents.operations.workbook import WorkbookOperation
from services.documents.worker import DocumentRequestError, DocumentWorkerPool
from tests.support.office_documents import (
    PNG_PIXEL,
    charted_workbook,
    memo_document,
    picture_presentation,
    report_presentation,
    rewrite_package,
    sales_workbook,
)

_ADAPTERS = {
    "pptx": TypeAdapter(list[PresentationOperation]),
    "xlsx": TypeAdapter(list[WorkbookOperation]),
    "docx": TypeAdapter(list[WordOperation]),
}


async def _edit(
    pool: DocumentWorkerPool,
    document_format: str,
    data: bytes,
    operations: list[dict[str, Any]],
    *,
    max_cells: int = 1_500_000,
    images: list[bytes] | None = None,
) -> tuple[dict[str, Any], bytes]:
    validated = _ADAPTERS[document_format].validate_python(operations)
    dumped = [item.model_dump(mode="json", exclude_none=True) for item in validated]
    for item in dumped:
        # The tool swaps image File references for attachment positions.
        if item.pop("image_file_id", None) is not None:
            item["image"] = 0
    result = await pool.run(
        "edit",
        document_format=document_format,
        data=data,
        attachments=images or [],
        args={
            "operations": dumped,
            "max_chars": 60_000,
            "source_ref": "file:source/revision:base",
            "author": "Report Agent",
            "max_cells": max_cells,
        },
    )
    assert result.data is not None
    return result.value, result.data


async def test_deck_replacement_across_runs_takes_the_first_runs_formatting(
    pool: DocumentWorkerPool,
) -> None:
    from pptx import Presentation

    _, output = await _edit(
        pool,
        "pptx",
        report_presentation(),
        [{"op": "replace_text", "find": "up 12", "replace": "rose 15"}],
    )

    slide = Presentation(io.BytesIO(output)).slides[0]
    runs = slide.placeholders[1].text_frame.paragraphs[0].runs
    assert [run.text for run in runs] == ["Revenue rose 15", "%"]
    assert runs[0].font.bold is None
    assert runs[1].font.bold is True
    assert runs[1].font.size.pt == 20


async def test_word_replacement_across_runs_takes_the_first_runs_formatting(
    pool: DocumentWorkerPool,
) -> None:
    from docx import Document

    _, output = await _edit(
        pool,
        "docx",
        memo_document(body_text="Plain "),
        [{"op": "replace_text", "find": "n bo", "replace": "n, BO"}],
    )

    paragraph = Document(io.BytesIO(output)).paragraphs[1]
    texts = [(run.text, run.bold) for run in paragraph.runs if run.text]
    assert texts == [("Plain, BO", None), ("ld", True)]


async def test_duplicate_slide_shares_images_and_refuses_charts(pool: DocumentWorkerPool) -> None:
    from pptx import Presentation

    data = picture_presentation()
    picture_id, chart_id = (slide.slide_id for slide in Presentation(io.BytesIO(data)).slides)

    result, output = await _edit(
        pool, "pptx", data, [{"op": "duplicate_slide", "slide_id": picture_id}]
    )

    slides = list(Presentation(io.BytesIO(output)).slides)
    copy = slides[1]
    assert copy.slide_id == result["changes"][0]["slide_id"]
    assert copy.shapes[0].image.blob == slides[0].shapes[0].image.blob
    with pytest.raises(DocumentRequestError, match=r"operations\[0\].*can't be copied"):
        await _edit(pool, "pptx", data, [{"op": "duplicate_slide", "slide_id": chart_id}])


async def test_workbook_with_a_chart_is_refused_rather_than_saved_without_it(
    pool: DocumentWorkerPool,
) -> None:
    with pytest.raises(DocumentRequestError, match=r"charts.*Sheets with drawings: Data"):
        await _edit(
            pool,
            "xlsx",
            charted_workbook(),
            [{"op": "set_cells", "sheet": "Notes", "anchor": "A1", "values": [["note"]]}],
        )


async def test_workbook_above_the_cell_limit_is_refused_before_loading(
    pool: DocumentWorkerPool,
) -> None:
    operations = [{"op": "set_cells", "sheet": "Sales", "anchor": "C1", "values": [[1]]}]

    def utf16(part: str, content: bytes) -> bytes:
        if not part.startswith("xl/worksheets/sheet"):
            return content
        declared = '<?xml version="1.0" encoding="UTF-16"?>' + content.decode()
        return declared.encode("utf-16")

    for data in (sales_workbook(rows=3), rewrite_package(sales_workbook(rows=3), utf16)):
        with pytest.raises(DocumentRequestError, match="more than the 8 cells that edits support"):
            await _edit(pool, "xlsx", data, operations, max_cells=8)


async def test_formula_naming_a_missing_sheet_fails(pool: DocumentWorkerPool) -> None:
    operations = [
        {"op": "set_cells", "sheet": "Sales", "anchor": "C2", "values": [["=B2*2"]]},
        {"op": "set_cells", "sheet": "Sales", "anchor": "C3", "values": [["=Budget!B3"]]},
    ]

    with pytest.raises(DocumentRequestError, match=r"operations\[1\]: Sales!C3 .*'Budget'"):
        await _edit(pool, "xlsx", sales_workbook(), operations)


async def test_sheet_references_follow_a_rename_and_block_deleting_a_used_sheet(
    pool: DocumentWorkerPool,
) -> None:
    from openpyxl import load_workbook

    data = sales_workbook()
    reference = [{"op": "set_cells", "sheet": "Notes", "anchor": "A1", "values": [["=Sales!B2"]]}]

    _, output = await _edit(
        pool,
        "xlsx",
        data,
        [*reference, {"op": "rename_sheet", "sheet": "Sales", "name": "Q1 Sales"}],
    )

    assert load_workbook(io.BytesIO(output))["Notes"]["A1"].value == "='Q1 Sales'!B2"
    with pytest.raises(DocumentRequestError, match=r"operations\[1\] \(delete_sheet\)"):
        await _edit(pool, "xlsx", data, [*reference, {"op": "delete_sheet", "sheet": "Sales"}])

    # openpyxl would number a title that differs only in case.
    result, output = await _edit(
        pool, "xlsx", data, [*reference, {"op": "rename_sheet", "sheet": "Sales", "name": "sales"}]
    )
    workbook = load_workbook(io.BytesIO(output))
    assert workbook.sheetnames == ["sales", "Notes"]
    assert workbook["Notes"]["A1"].value == "=sales!B2"
    assert result["changes"][1]["sheet"] == "sales"


async def test_table_headers_must_stay_in_step_with_their_columns(
    pool: DocumentWorkerPool,
) -> None:
    from openpyxl.worksheet.table import Table

    def build(workbook: Any) -> None:
        sheet = workbook.active
        sheet.title = "Data"
        for row in (["Amount"], [5]):
            sheet.append(row)
        sheet["C1"] = "=SUM(Revenue[Amount])"
        sheet.add_table(Table(displayName="Revenue", ref="A1:A2"))

    rename_column = [{"op": "set_cells", "sheet": "Data", "anchor": "A1", "values": [["Total"]]}]
    with pytest.raises(DocumentRequestError, match=r"operations\[0\].*column name.*'Revenue'"):
        await _edit(pool, "xlsx", _workbook(build), rename_column)

    new_table = [
        {"op": "set_cells", "sheet": "Data", "anchor": "E1", "values": [["Region"], ["North"]]},
        {"op": "add_table", "sheet": "Data", "name": "Regions", "range": "E1:E2"},
        {"op": "set_cells", "sheet": "Data", "anchor": "E1", "values": [[None]]},
    ]
    with pytest.raises(DocumentRequestError, match=r"operations\[2\].*needs text.*'Regions'"):
        await _edit(pool, "xlsx", _workbook(build), new_table)


async def test_workbook_without_calculation_properties_saves_with_recalculation(
    pool: DocumentWorkerPool,
) -> None:
    import re

    from openpyxl import load_workbook

    def drop_calculation(part: str, content: bytes) -> bytes:
        return re.sub(rb"<calcPr[^>]*/>", b"", content) if part == "xl/workbook.xml" else content

    data = rewrite_package(sales_workbook(), drop_calculation)
    assert load_workbook(io.BytesIO(data)).calculation is None

    _, output = await _edit(
        pool, "xlsx", data, [{"op": "set_cells", "sheet": "Notes", "anchor": "A1", "values": [[1]]}]
    )

    workbook = load_workbook(io.BytesIO(output))
    assert (workbook["Notes"]["A1"].value, workbook.calculation.fullCalcOnLoad) == (1, True)


async def test_dropdown_rejects_values_outside_its_list(pool: DocumentWorkerPool) -> None:
    from openpyxl import load_workbook

    operation = {
        "op": "add_data_validation",
        "sheet": "Notes",
        "range": "A1:A5",
        "options": ["Open", "Closed"],
    }

    _, output = await _edit(pool, "xlsx", sales_workbook(), [operation])

    [validation] = load_workbook(io.BytesIO(output))["Notes"].data_validations.dataValidation
    assert (validation.showErrorMessage, validation.errorStyle) == (True, "stop")


async def test_freeze_panes_outside_the_grid_is_refused(pool: DocumentWorkerPool) -> None:
    from openpyxl import load_workbook

    freeze = {"op": "freeze_panes", "sheet": "Notes", "cell": "XFE1"}
    with pytest.raises(DocumentRequestError, match="outside the sheet"):
        await _edit(pool, "xlsx", sales_workbook(), [freeze])

    _, output = await _edit(pool, "xlsx", sales_workbook(), [{**freeze, "cell": "XFD2"}])
    assert load_workbook(io.BytesIO(output))["Notes"].freeze_panes == "XFD2"


async def test_word_edit_refuses_a_stale_paragraph_without_quoting_it(
    pool: DocumentWorkerPool,
) -> None:
    secret = "Confidential body text"
    operations = [{"op": "set_paragraph", "index": 1, "expect_text": "Other", "runs": ["New"]}]

    with pytest.raises(DocumentRequestError) as refused:
        await _edit(pool, "docx", memo_document(body_text=secret), operations)

    assert "doesn't start with expect_text" in refused.value.message
    assert secret not in refused.value.message


def _saved(document: Any) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _workbook(build: Callable[[Any], None]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    build(workbook)
    return _saved(workbook)


async def test_workbook_with_parts_openpyxl_drops_is_refused(pool: DocumentWorkerPool) -> None:
    relationship = (
        '<Relationship Id="rIdCustom" Target="../customXml/item1.xml" Type="http://schemas.'
        'openxmlformats.org/officeDocument/2006/relationships/customXml"/></Relationships>'
    )

    def link(part: str, content: bytes) -> bytes:
        if part != "xl/_rels/workbook.xml.rels":
            return content
        return content.replace(b"</Relationships>", relationship.encode())

    data = rewrite_package(
        sales_workbook(), link, extra={"customXml/item1.xml": b"<data>Keep me</data>"}
    )
    operations = [{"op": "set_cells", "sheet": "Notes", "anchor": "A1", "values": [["note"]]}]

    with pytest.raises(DocumentRequestError, match="custom XML data, which editing would remove"):
        await _edit(pool, "xlsx", data, operations)


async def test_merge_over_a_populated_cell_is_refused(pool: DocumentWorkerPool) -> None:
    from openpyxl import load_workbook

    def build(workbook: Any) -> None:
        workbook.active.title = "Data"
        workbook.active.append(["keep", "lost"])

    merge = {"op": "merge_cells", "sheet": "Data", "range": "A1:B1"}
    with pytest.raises(DocumentRequestError, match=r"operations\[0\].*B1 has content"):
        await _edit(pool, "xlsx", _workbook(build), [merge])

    _, output = await _edit(pool, "xlsx", _workbook(build), [{**merge, "range": "A2:B2"}])
    assert "A2:B2" in load_workbook(io.BytesIO(output))["Data"].merged_cells


async def test_append_refuses_a_table_ending_in_a_totals_row(pool: DocumentWorkerPool) -> None:
    from openpyxl.worksheet.table import Table

    def build(workbook: Any) -> None:
        sheet = workbook.active
        sheet.title = "Data"
        for row in (["Amount"], [5], ["=SUBTOTAL(109,[Amount])"]):
            sheet.append(row)
        sheet.add_table(Table(displayName="Totals", ref="A1:A3", totalsRowCount=1))

    with pytest.raises(DocumentRequestError, match="totals row"):
        await _edit(
            pool, "xlsx", _workbook(build), [{"op": "append_rows", "sheet": "Data", "rows": [[7]]}]
        )


async def test_append_extends_a_plain_table_and_its_filter(pool: DocumentWorkerPool) -> None:
    from openpyxl import load_workbook
    from openpyxl.worksheet.table import Table

    table_name = "T" * 1_500

    def build(workbook: Any) -> None:
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Amount"])
        sheet.append([5])
        sheet.add_table(Table(displayName=table_name, ref="A1:A2"))

    result, output = await _edit(
        pool, "xlsx", _workbook(build), [{"op": "append_rows", "sheet": "Data", "rows": [[7]]}]
    )

    table = next(iter(load_workbook(io.BytesIO(output))["Data"].tables.values()))
    assert (table.ref, table.autoFilter.ref) == ("A1:A3", "A1:A3")
    # Uploaded table names stay within the capped identifier contract.
    change = result["changes"][0]
    assert table_name not in change["summary"]
    assert change["tables"] == [table_name[:100]]


async def test_append_writes_under_a_table_that_starts_past_column_a(
    pool: DocumentWorkerPool,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.worksheet.table import Table

    def build(workbook: Any) -> None:
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append([None, "Region", "Amount"])
        sheet.append([None, "North", 5])
        sheet.add_table(Table(displayName="Sales", ref="B1:C2"))

    _, output = await _edit(
        pool,
        "xlsx",
        _workbook(build),
        [{"op": "append_rows", "sheet": "Data", "rows": [["South", 7]]}],
    )

    sheet = load_workbook(io.BytesIO(output))["Data"]
    assert [cell.value for cell in sheet[3]] == [None, "South", 7]
    assert sheet.tables["Sales"].ref == "B1:C3"


async def test_deleting_a_sheet_a_surviving_name_uses_is_refused(pool: DocumentWorkerPool) -> None:
    from openpyxl.workbook.defined_name import DefinedName

    def build(workbook: Any) -> None:
        workbook.active.title = "Data"
        notes = workbook.create_sheet("Notes")
        workbook.defined_names["Total"] = DefinedName("Total", attr_text="Data!$A$1")
        workbook["Data"].defined_names["Local"] = DefinedName("Local", attr_text="Data!$A$2")
        notes["A1"] = "=Total"

    delete = [{"op": "delete_sheet", "sheet": "Data"}]
    with pytest.raises(DocumentRequestError, match=r"delete_sheet.*the name Total still refers"):
        await _edit(pool, "xlsx", _workbook(build), delete)

    # A name local to the deleted sheet goes with it.
    def local_only(workbook: Any) -> None:
        workbook.active.title = "Data"
        workbook.create_sheet("Notes")
        workbook["Data"].defined_names["Local"] = DefinedName("Local", attr_text="Data!$A$2")

    await _edit(pool, "xlsx", _workbook(local_only), delete)


async def test_names_are_checked_as_they_end_and_from_their_sheet(
    pool: DocumentWorkerPool,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.workbook.defined_name import DefinedName

    data = sales_workbook()
    _, output = await _edit(
        pool,
        "xlsx",
        data,
        [
            {"op": "define_name", "name": "Revenue", "refers_to": "Sales!$B$2"},
            {"op": "rename_sheet", "sheet": "Sales", "name": "Q1"},
        ],
    )
    assert load_workbook(io.BytesIO(output)).defined_names["Revenue"].attr_text == "'Q1'!$B$2"

    def local(workbook: Any) -> None:
        workbook.active.title = "Sales"
        workbook.create_sheet("Notes")
        workbook["Sales"].defined_names["LocalOnly"] = DefinedName(
            "LocalOnly", attr_text="Sales!$A$1"
        )

    def write(sheet: str, formula: str) -> list[dict[str, Any]]:
        return [{"op": "set_cells", "sheet": sheet, "anchor": "B1", "values": [[formula]]}]

    with pytest.raises(DocumentRequestError, match="Notes!B1 uses the name 'LocalOnly'"):
        await _edit(pool, "xlsx", _workbook(local), write("Notes", "=LocalOnly"))
    await _edit(pool, "xlsx", _workbook(local), write("Sales", "=LocalOnly"))
    await _edit(pool, "xlsx", _workbook(local), write("Notes", "=Sales!LocalOnly"))


async def test_chart_added_before_a_rename_follows_the_sheet(pool: DocumentWorkerPool) -> None:
    import zipfile

    chart = {
        "op": "add_chart",
        "sheet": "Sales",
        "anchor": "F2",
        "chart_type": "column",
        "data": "B1:B4",
        "categories": "A2:A4",
    }
    _, output = await _edit(
        pool,
        "xlsx",
        sales_workbook(),
        [chart, {"op": "rename_sheet", "sheet": "Sales", "name": "Q1"}],
    )

    with zipfile.ZipFile(io.BytesIO(output)) as archive:
        saved = next(
            archive.read(part).decode()
            for part in archive.namelist()
            if part.startswith("xl/charts/")
        )
    assert "'Q1'!$B$2:$B$4" in saved
    assert "Sales" not in saved
    # A chart removed with its sheet no longer refers to anything.
    await _edit(pool, "xlsx", sales_workbook(), [chart, {"op": "delete_sheet", "sheet": "Sales"}])


async def test_table_column_formulas_follow_a_rename(pool: DocumentWorkerPool) -> None:
    from openpyxl import load_workbook
    from openpyxl.worksheet.table import Table, TableFormula

    def build(workbook: Any) -> None:
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Amount", "Scaled"])
        sheet.append([5, "=Rates!$A$1*2"])
        table = Table(displayName="Scaled", ref="A1:B2")
        table._initialise_columns()
        table.tableColumns[1].calculatedColumnFormula = TableFormula(attr_text="Rates!$A$1*2")
        sheet.add_table(table)
        workbook.create_sheet("Rates")["A1"] = 3

    _, output = await _edit(
        pool,
        "xlsx",
        _workbook(build),
        [{"op": "rename_sheet", "sheet": "Rates", "name": "Prices"}],
    )

    table = load_workbook(io.BytesIO(output))["Data"].tables["Scaled"]
    assert table.tableColumns[1].calculatedColumnFormula.attr_text == "Prices!$A$1*2"


async def test_workbook_readback_reports_cleared_cells(pool: DocumentWorkerPool) -> None:
    result, _ = await _edit(
        pool,
        "xlsx",
        sales_workbook(),
        [{"op": "set_cells", "sheet": "Sales", "anchor": "A2", "values": [[None]]}],
    )

    assert result["readback"]["blocks"][0]["cells"] == [{"ref": "A2", "value": None}]


async def test_word_edit_of_an_inherited_header_keeps_its_field(pool: DocumentWorkerPool) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    document.add_paragraph("First section")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    document.sections[0].header.paragraphs[0]._p.append(field)
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Second section")

    _, output = await _edit(
        pool,
        "docx",
        _saved(document),
        [{"op": "set_header", "section": 2, "text": "Appendix"}],
    )

    sections = Document(io.BytesIO(output)).sections
    fields = [
        len(section.header._element.findall(f".//{qn('w:fldSimple')}")) for section in sections
    ]
    assert fields == [1, 1]
    assert "Appendix" in sections[1].header._element.xpath("string(.)")
    assert "Appendix" not in sections[0].header._element.xpath("string(.)")


async def test_cleared_word_template_keeps_its_sections(pool: DocumentWorkerPool) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION

    document = Document()
    document.add_paragraph("Cover")
    document.add_section(WD_SECTION.NEW_PAGE).orientation = WD_ORIENT.LANDSCAPE
    document.add_paragraph("Body")
    pool_args = [{"op": "insert_paragraphs", "paragraphs": ["New body"]}]

    result = await pool.run(
        "edit",
        document_format="docx",
        data=_saved(document),
        args={
            "operations": pool_args,
            "clear": True,
            "max_chars": 60_000,
            "source_ref": "template:default.docx",
            "author": "Report Agent",
            "max_cells": 1_500_000,
        },
    )

    assert result.data is not None
    created = Document(io.BytesIO(result.data))
    assert [section.orientation for section in created.sections] == [
        WD_ORIENT.PORTRAIT,
        WD_ORIENT.LANDSCAPE,
    ]
    assert [paragraph.text for paragraph in created.paragraphs] == ["", "New body"]


async def test_word_write_through_a_merged_cell_is_refused(pool: DocumentWorkerPool) -> None:
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    data = _saved(document)

    with pytest.raises(DocumentRequestError, match="which a merged cell covers"):
        await _edit(
            pool,
            "docx",
            data,
            [{"op": "set_table_cells", "table_index": 0, "values": [["first", "second"]]}],
        )

    _, output = await _edit(
        pool,
        "docx",
        data,
        [{"op": "set_table_cells", "table_index": 0, "values": [["merged", ""], ["a", "b"]]}],
    )
    rows = Document(io.BytesIO(output)).tables[0].rows
    assert [cell.text for cell in rows[0].cells] == ["merged", "merged"]
    assert [cell.text for cell in rows[1].cells] == ["a", "b"]


async def test_word_table_rows_copy_formatting_without_merges_or_bookmarks(
    pool: DocumentWorkerPool,
) -> None:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Pt

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0))
    table.rows[1].height = Pt(30)
    table.cell(1, 1).paragraphs[0]._p.append(
        parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="1" w:name="Target"/>')
    )
    operation = {"op": "set_table_cells", "table_index": 0, "start_row": 2, "values": [["", "new"]]}

    _, output = await _edit(pool, "docx", _saved(document), [operation])

    saved = Document(io.BytesIO(output)).tables[0]
    assert saved.cell(2, 0)._tc is not saved.cell(1, 0)._tc
    assert [cell.text for cell in saved.rows[2].cells] == ["", "new"]
    assert saved.rows[2].height == Pt(30)
    assert len(saved._tbl.findall(f".//{qn('w:bookmarkStart')}")) == 1


async def test_word_text_rewrites_keep_bookmarks_around_their_text(
    pool: DocumentWorkerPool,
) -> None:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    def bookmarked(*parts: str) -> bytes:
        document = Document()
        paragraph = document.add_paragraph()
        for part in parts:
            if part in ("[", "]"):
                tag = "bookmarkStart" if part == "[" else "bookmarkEnd"
                name = ' w:name="Target"' if part == "[" else ""
                paragraph._p.append(parse_xml(f'<w:{tag} {nsdecls("w")} w:id="1"{name}/>'))
            else:
                paragraph.add_run(part)
        return _saved(document)

    def rewrite(expect_text: str) -> list[dict[str, Any]]:
        return [{"op": "set_paragraph", "index": 0, "expect_text": expect_text, "runs": ["New"]}]

    _, output = await _edit(pool, "docx", bookmarked("[", "Old", "]"), rewrite("Old"))
    children = Document(io.BytesIO(output)).paragraphs[0]._p
    assert [child.tag.rsplit("}", 1)[1] for child in children] == [
        "bookmarkStart",
        "r",
        "bookmarkEnd",
    ]
    with pytest.raises(DocumentRequestError, match="bookmark covers part of this text"):
        await _edit(pool, "docx", bookmarked("Keep ", "[", "this", "]"), rewrite("Keep"))


async def test_deleting_paragraphs_keeps_comment_and_bookmark_ranges_whole(
    pool: DocumentWorkerPool,
) -> None:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = Document()
    first, second = document.add_paragraph("First"), document.add_paragraph("Second")
    document.add_comment((first.runs[0], second.runs[0]), text="Spans both")
    document.add_paragraph("Third")._p.append(
        parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="7" w:name="Target"/>')
    )
    document.add_paragraph("Fourth")._p.append(
        parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="7"/>')
    )
    data = _saved(document)

    def delete(index: int, count: int, expect_text: str) -> list[dict[str, Any]]:
        return [
            {"op": "delete_paragraphs", "index": index, "count": count, "expect_text": expect_text}
        ]

    for index, expect_text, kind in ((0, "First", "comment"), (2, "Third", "bookmark")):
        with pytest.raises(DocumentRequestError, match=f"A {kind} covers text inside and outside"):
            await _edit(pool, "docx", data, delete(index, 1, expect_text))

    _, output = await _edit(pool, "docx", data, delete(0, 2, "First"))
    assert list(Document(io.BytesIO(output)).comments) == []


async def test_word_image_fits_the_section_it_is_inserted_in(pool: DocumentWorkerPool) -> None:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.add_paragraph("Narrow")
    document.sections[0].page_width = Pt(360)
    document.add_section().page_width = Pt(1000)
    document.add_paragraph("Wide")
    operation = {
        "op": "add_image",
        "image_file_id": {"entity_id": str(uuid4()), "label": "logo.png"},
        "width": 500,
        "index": 0,
        "expect_text": "Narrow",
        "position": "after",
    }

    _, output = await _edit(pool, "docx", _saved(document), [operation], images=[PNG_PIXEL])

    [picture] = Document(io.BytesIO(output)).inline_shapes
    room = Pt(360) - document.sections[0].left_margin - document.sections[0].right_margin
    assert (picture.width, picture.height) == (room, room)


async def test_word_readback_covers_paragraphs_headers_comments_and_controls(
    pool: DocumentWorkerPool,
) -> None:
    from docx import Document
    from docx.oxml import OxmlElement

    document = Document(io.BytesIO(memo_document()))
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    content.append(Document().add_paragraph("Client name")._p)
    control.append(content)
    document.element.body.sectPr.addprevious(control)

    result, _ = await _edit(
        pool,
        "docx",
        _saved(document),
        [
            {"op": "set_paragraph", "index": 0, "expect_text": "Title", "runs": ["Renamed"]},
            {"op": "set_header", "text": "New header"},
            {"op": "add_comment", "index": 1, "expect_text": "Plain", "text": "Agent note"},
            {"op": "replace_text", "find": "Client name", "replace": "Acme"},
        ],
    )

    readback = result["readback"]
    texts = [block["text"]["content"] for block in readback["blocks"]]
    assert texts == ["Renamed", "Plain bold", "Acme"]
    assert readback["headers"][0]["text"]["content"] == "New header"
    assert readback["comments"][0]["text"]["content"] == "Agent note"


async def test_word_replacement_keeps_a_nonbreaking_hyphen(pool: DocumentWorkerPool) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    run = document.add_paragraph().add_run("A")
    run._r.append(OxmlElement("w:noBreakHyphen"))
    text = OxmlElement("w:t")
    text.text = "B"
    run._r.append(text)
    document.add_paragraph("A later")

    _, output = await _edit(
        pool, "docx", _saved(document), [{"op": "replace_text", "find": "A", "replace": "C"}]
    )

    paragraphs = Document(io.BytesIO(output)).paragraphs
    assert paragraphs[0]._p.findall(f".//{qn('w:noBreakHyphen')}")
    assert paragraphs[1].text == "C later"


async def test_resizing_a_deck_table_through_a_vertical_merge_is_refused(
    pool: DocumentWorkerPool,
) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_table(2, 1, Inches(1), Inches(1), Inches(3), Inches(1))
    frame.table.cell(0, 0).merge(frame.table.cell(1, 0))
    operation = {
        "op": "set_table",
        "slide_id": slide.slide_id,
        "shape_id": frame.shape_id,
        "rows": [["Only"]],
    }

    with pytest.raises(DocumentRequestError, match="cut through a merged cell"):
        await _edit(pool, "pptx", _saved(presentation), [operation])


async def test_duplicate_slide_keeps_its_transition_and_notes_formatting(
    pool: DocumentWorkerPool,
) -> None:
    from pptx import Presentation
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide._element.append(parse_xml(f"<p:transition {nsdecls('p')}><p:fade/></p:transition>"))
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Say "
    emphasis = notes.paragraphs[0].add_run()
    emphasis.text = "this"
    emphasis.font.bold = True

    result, output = await _edit(
        pool, "pptx", _saved(presentation), [{"op": "duplicate_slide", "slide_id": slide.slide_id}]
    )

    copy = Presentation(io.BytesIO(output)).slides.get(result["changes"][0]["slide_id"])
    assert copy._element.xpath("./p:transition/p:fade")
    runs = copy.notes_slide.notes_text_frame.paragraphs[0].runs
    assert [(run.text, run.font.bold) for run in runs] == [("Say ", None), ("this", True)]


async def test_deleted_slide_leaves_the_deck_sections(pool: DocumentWorkerPool) -> None:
    from lxml import etree
    from pptx import Presentation

    sections = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    presentation = Presentation()
    kept, deleted = (
        presentation.slides.add_slide(presentation.slide_layouts[6]).slide_id for _ in range(2)
    )
    presentation._element.append(
        etree.fromstring(
            '<p:extLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:ext uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}">'
            f'<p14:sectionLst xmlns:p14="{sections}"><p14:section name="All" '
            'id="{2B1D4E1A-8C2E-4E57-9F0B-6C1C1E1D2A3B}"><p14:sldIdLst>'
            f'<p14:sldId id="{kept}"/><p14:sldId id="{deleted}"/>'
            "</p14:sldIdLst></p14:section></p14:sectionLst></p:ext></p:extLst>"
        )
    )

    _, output = await _edit(
        pool, "pptx", _saved(presentation), [{"op": "delete_slide", "slide_id": deleted}]
    )

    saved = Presentation(io.BytesIO(output))._element
    listed = [entry.get("id") for entry in saved.iter(f"{{{sections}}}sldId")]
    assert listed == [str(kept)]


async def test_deck_edit_read_back_leaves_inherited_text_colour_alone(
    pool: DocumentWorkerPool,
) -> None:
    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[5].name

    _, output = await _edit(
        pool,
        "pptx",
        _saved(presentation),
        [{"op": "add_slide", "layout": layout, "placeholders": [{"idx": 0, "paragraphs": ["Q3"]}]}],
    )

    title = Presentation(io.BytesIO(output)).slides[0].shapes.title
    [run] = title.text_frame.paragraphs[0].runs
    # An empty solid fill renders black, hiding text on a dark layout.
    assert run.font.fill.type is None


async def test_new_deck_table_takes_the_decks_default_table_style(
    pool: DocumentWorkerPool,
) -> None:
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE

    style_id = "{6F3C2E8A-5B1D-4C8E-9A7F-3D2B1E0C4A51}"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    styles = presentation.part.part_related_by(RELATIONSHIP_TYPE.TABLE_STYLES)
    styles._blob = (
        '<a:tblStyleLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'def="{style_id}"/>'
    ).encode()

    result, output = await _edit(
        pool,
        "pptx",
        _saved(presentation),
        [
            {
                "op": "add_table",
                "slide_id": slide.slide_id,
                "rows": [["Region"], ["EMEA"]],
                "left": 10,
                "top": 10,
                "width": 200,
                "height": 60,
            }
        ],
    )

    saved = Presentation(io.BytesIO(output)).slides[0]
    [table] = [
        shape for shape in saved.shapes if shape.shape_id == result["changes"][0]["shape_id"]
    ]
    assert table._element.xpath(".//a:tableStyleId/text()") == [style_id]


async def test_large_edit_result_stays_within_the_page_budget(pool: DocumentWorkerPool) -> None:
    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[6]
    layout.name = "L" * 100
    operations = [{"op": "add_slide", "layout": layout.name}] * 500

    result, _ = await _edit(pool, "pptx", _saved(presentation), operations)

    # The rest of the page holds the saved File's fields.
    assert len(json.dumps(result, separators=(",", ":"))) <= 60_000 - 1_000
    assert len(result["changes"]) + result["changes_omitted"] == 500
    assert result["readback"]["slides_truncated"] is True


def test_warnings_past_the_cap_are_counted_as_omitted() -> None:
    from services.documents.editing import EditLog

    log = EditLog({"max_chars": 60_000, "source_ref": "file:source/revision:base"})
    for number in range(51):
        log.warn(f"Warning {number}.")

    result = log.result(dict)

    assert (len(result["warnings"]), result["warnings_omitted"]) == (50, 1)


def test_a_failed_allocation_during_an_edit_reaches_the_worker_loop() -> None:
    from services.documents.editing import EditLog, run_operations

    def allocate(_operation: dict[str, Any], _log: EditLog) -> None:
        raise MemoryError

    log = EditLog({"max_chars": 60_000, "source_ref": "file:source/revision:base"})
    # The worker exits on MemoryError, so the edit loop mustn't turn it into a request error.
    with pytest.raises(MemoryError):
        run_operations([{"op": "allocate"}], {"allocate": allocate}, log)
