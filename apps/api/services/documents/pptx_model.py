# apps/api/services/documents/pptx_model.py

"""Bounded, paged reads of presentations. Runs only inside the document worker."""

import math
from typing import Any

from services.documents.reading import (
    ReadPage,
    color,
    compact,
    external_links,
    formatting,
    name,
    points,
    spans,
)

_MAX_GROUP_DEPTH = 5
# Every layout repeats these, and edits don't target them.
_LAYOUT_FURNITURE = frozenset({"date", "footer", "slide_number"})


def read_presentation(presentation: Any, args: dict[str, Any]) -> dict[str, Any]:
    """Returns as many requested slides as fit the page.

    The first page, when `slides` is omitted, also lists layouts and external links.
    """
    page = ReadPage(args)
    slides = list(presentation.slides)
    requested = args.get("slides")
    result: dict[str, Any] = {
        "slide_count": len(slides),
        "slide_size": {
            "width": points(presentation.slide_width),
            "height": points(presentation.slide_height),
        },
    }
    if requested is None:
        result.update(_layouts(presentation, page))
        result.update(external_links(_external_relationships(presentation), page))
    result["slides"] = []
    missing = []
    for number in requested or range(1, len(slides) + 1):
        if not 1 <= number <= len(slides):
            missing.append(number)
            continue
        item = read_slide(slides[number - 1], number, page)
        too_large = f"Slide {number} is larger than one read page. Read the other slides by number."
        if not page.add(item, first=not result["slides"], too_large=too_large):
            result["next_slide"] = number
            break
        result["slides"].append(item)
    if missing:
        result["missing_slides"] = missing
    return result


def image_parts(element: Any, part: Any) -> list[str]:
    """Returns part names of images embedded in an element, resolved through its part."""
    refs = []
    for rel_id in element.xpath(".//a:blip/@r:embed"):
        try:
            ref = str(part.related_part(rel_id).partname).lstrip("/")
        except KeyError:
            continue
        if ref not in refs:
            refs.append(ref)
    return refs


def _layouts(presentation: Any, page: ReadPage) -> dict[str, Any]:
    layouts = []
    for master in presentation.slide_masters:
        for layout in master.slide_layouts:
            placeholders = [_placeholder(shape) for shape in layout.placeholders]
            item = {
                "name": name(layout.name),
                "placeholders": [
                    placeholder
                    for placeholder in placeholders
                    if placeholder["type"] not in _LAYOUT_FURNITURE
                ],
            }
            if not page.fits(item):
                return {"layouts": layouts, "layouts_truncated": True}
            layouts.append(item)
    return {"layouts": layouts}


def _external_relationships(presentation: Any) -> Any:
    for relationship in presentation.part.package.iter_rels():
        if relationship.is_external:
            yield relationship.reltype, relationship.target_ref


def read_slide(slide: Any, number: int, page: ReadPage) -> dict[str, Any]:
    """Returns one slide's layout, shapes, and notes as reads show them."""
    notes = ""
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        notes = slide.notes_slide.notes_text_frame.text
    background = slide._element.xpath("./p:cSld/p:bg")
    return compact(
        {
            "number": number,
            "slide_id": slide.slide_id,
            "layout": name(slide.slide_layout.name),
            "background_images": image_parts(background[0], slide.part) if background else None,
            "shapes": [_shape(shape, page, depth=0) for shape in slide.shapes],
            "notes": page.text(notes),
        }
    )


def _shape(shape: Any, page: ReadPage, *, depth: int) -> dict[str, Any]:
    item: dict[str, Any] = {
        "shape_id": shape.shape_id,
        "name": name(shape.name),
        "kind": _kind(shape),
        "placeholder": _placeholder(shape) if shape.is_placeholder else None,
        "position": compact(
            {
                "left": points(shape.left),
                "top": points(shape.top),
                "width": points(shape.width),
                "height": points(shape.height),
            }
        ),
    }
    if shape.has_text_frame:
        item["paragraphs"] = [
            _paragraph(paragraph, page) for paragraph in shape.text_frame.paragraphs
        ]
    if getattr(shape, "has_table", False):
        item["table"] = [[page.text(cell.text) for cell in row.cells] for row in shape.table.rows]
    if getattr(shape, "has_chart", False):
        item["chart"] = _chart(shape.chart, page)
    if item["kind"] != "group":
        item["images"] = image_parts(shape._element, shape.part)
    elif depth < _MAX_GROUP_DEPTH:
        item["shapes"] = [_shape(child, page, depth=depth + 1) for child in shape.shapes]
    else:
        # Deeper groups stay unread so a hostile file can't nest without bound.
        item["children_omitted"] = len(shape.shapes)
    return compact(item)


def _kind(shape: Any) -> str:
    try:
        shape_type = shape.shape_type
    except NotImplementedError:
        return "other"
    if shape_type is None:
        return "other"
    kind = shape_type.name.lower()
    # Picture placeholders report as placeholders; the image is what matters.
    if kind == "placeholder" and shape._element.xpath("./p:blipFill"):
        return "picture"
    return kind


def _placeholder(shape: Any) -> dict[str, Any]:
    placeholder = shape.placeholder_format
    return {
        "idx": placeholder.idx,
        "type": placeholder.type.name.lower() if placeholder.type is not None else None,
        "name": name(shape.name),
    }


def _paragraph(paragraph: Any, page: ReadPage) -> dict[str, Any]:
    text, ranges = spans(_pieces(paragraph))
    return {
        "text": page.text(text),
        **compact({"level": paragraph.level or None, "spans": ranges}),
    }


def _pieces(paragraph: Any) -> Any:
    """Yields runs, fields, and soft line breaks in order, with breaks as `\\v` like python-pptx."""
    from pptx.text.text import _Run

    for element in paragraph._element.content_children:
        if element.tag.endswith("}r"):
            yield element.text, _run_style(_Run(element, paragraph).font)
        else:
            yield element.text, {}


def _run_style(font: Any) -> dict[str, Any]:
    return formatting(
        {
            "bold": font.bold,
            "italic": font.italic,
            "underline": bool(font.underline) if font.underline is not None else None,
            "size": font.size.pt if font.size is not None else None,
            "font": name(font.name) if font.name else None,
            "color": color(font.color),
        }
    )


def _chart(chart: Any, page: ReadPage) -> dict[str, Any]:
    """Returns categories and series; XY and bubble series also carry x values and sizes."""
    try:
        chart_type = chart.chart_type.name.lower()
    except (NotImplementedError, ValueError, KeyError):
        chart_type = "other"
    plots = list(chart.plots)
    categories = list(plots[0].categories) if plots else []
    return {
        "type": chart_type,
        "categories": [page.value(category) for category in categories],
        "series": [
            compact(
                {
                    "name": page.text(series.name or ""),
                    "plot": plot_index if len(plots) > 1 else None,
                    "values": list(series.values),
                    "x": _point_values(series, "xVal", page),
                    "sizes": _point_values(series, "bubbleSize", page),
                }
            )
            for plot_index, plot in enumerate(plots)
            for series in plot.series
        ],
    }


def _point_values(series: Any, element: str, page: ReadPage) -> list[Any]:
    """Returns a series' cached points for `c:xVal` or `c:bubbleSize`, in index order."""
    points_by_index = {}
    for point in series._element.xpath(f"./c:{element}//c:pt"):
        values = point.xpath("./c:v/text()")
        points_by_index[int(point.get("idx", 0))] = _number(values[0] if values else "", page)
    return [points_by_index[index] for index in sorted(points_by_index)]


def _number(text: str, page: ReadPage) -> Any:
    try:
        number = float(text)
    except ValueError:
        return page.text(text)
    return number if math.isfinite(number) else page.text(text)
