# apps/api/services/documents/docx_edit.py

"""Edits to Word documents. Runs only inside the document worker.

Paragraph and table indexes count only paragraphs and tables directly in the
body, as reads number them. Operations that would drop images, fields,
comment anchors, or section breaks fail instead, as do rewrites and deletions
that would leave a bookmark or comment covering the wrong text.
"""

import io
from collections.abc import Iterable, Iterator
from copy import deepcopy
from typing import Any

from services.documents.docx_model import has_revisions, read_block, read_comment, read_story
from services.documents.editing import EditLog, OperationError, replace_in_runs, run_operations

_WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_SIMPLE_RUN_PARTS = frozenset(
    f"{_WORD}{tag}"
    for tag in ("rPr", "t", "tab", "cr", "noBreakHyphen", "softHyphen", "lastRenderedPageBreak")
)
# Hyphens that change line breaking; rewriting run text would turn them into plain hyphens.
_SEMANTIC_RUN_PARTS = frozenset(f"{_WORD}{tag}" for tag in ("noBreakHyphen", "softHyphen"))
_SIMPLE_PARAGRAPH_PARTS = frozenset(
    f"{_WORD}{tag}" for tag in ("pPr", "bookmarkStart", "bookmarkEnd", "proofErr")
)
# Markers between runs that text replacement can move text across.
_MARKERS = frozenset(
    f"{_WORD}{tag}"
    for tag in ("bookmarkStart", "bookmarkEnd", "commentRangeStart", "commentRangeEnd", "proofErr")
)
_STORIES = {
    ("set_header", "primary"): "header",
    ("set_header", "first_page"): "first_page_header",
    ("set_header", "even_page"): "even_page_header",
    ("set_footer", "primary"): "footer",
    ("set_footer", "first_page"): "first_page_footer",
    ("set_footer", "even_page"): "even_page_footer",
}
_BLOCKS = frozenset(f"{_WORD}{tag}" for tag in ("p", "tbl", "sdt"))
# Range endpoints by the kind of range, which each have a w:id shared by their ends.
_RANGE_ENDS = {
    f"{_WORD}bookmarkStart": "bookmark",
    f"{_WORD}bookmarkEnd": "bookmark",
    f"{_WORD}commentRangeStart": "comment",
    f"{_WORD}commentRangeEnd": "comment",
    f"{_WORD}commentReference": "comment",
}
_BOOKMARK_ENDS = frozenset((f"{_WORD}bookmarkStart", f"{_WORD}bookmarkEnd"))
# Row properties a new row doesn't take from the row it copies.
_UNCOPIED_ROW_PARTS = frozenset(
    f"{_WORD}{tag}" for tag in ("tblHeader", "ins", "del", "trPrChange")
)
_R_NAMESPACE = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_COMMENTS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
_IMAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


def edit_document(document: Any, args: dict[str, Any], images: list[bytes]) -> dict[str, Any]:
    """Applies operations to an open document and returns changes, warnings, and read-back.

    With `clear`, the body and comments are emptied first, keeping styles,
    headers, footers, and sections. Each section break stays as an empty paragraph.
    """
    log = EditLog(args)
    editor = _DocumentEditor(document, images, str(args.get("author") or "Agent"))
    if args.get("clear"):
        editor.clear()
    elif has_revisions(document):
        log.warn("The document has tracked changes. These edits aren't tracked.")
    run_operations(args["operations"], editor.handlers(), log)
    editor.end_with_paragraph()
    return log.result(lambda: editor.readback(log))


class _DocumentEditor:
    def __init__(self, document: Any, images: list[bytes], author: str) -> None:
        self.document = document
        self.body = document.element.body
        self.images = images
        self.author = author
        self.touched: dict[Any, None] = {}
        # (operation, section number, variant) of headers and footers set, for the read-back.
        self.stories: dict[tuple[str, int, str], None] = {}
        self.comment_ids: set[int] = set()

    def handlers(self) -> dict[str, Any]:
        return {
            "replace_text": self.replace_text,
            "insert_paragraphs": self.insert_paragraphs,
            "set_paragraph": self.set_paragraph,
            "delete_paragraphs": self.delete_paragraphs,
            "insert_table": self.insert_table,
            "set_table_cells": self.set_table_cells,
            "add_image": self.add_image,
            "page_break": self.page_break,
            "set_header": self.set_story,
            "set_footer": self.set_story,
            "add_comment": self.add_comment,
        }

    def clear(self) -> None:
        for child in list(self.body):
            if child.tag == f"{_WORD}sectPr":
                continue
            properties = child.find(f"{_WORD}pPr") if child.tag == f"{_WORD}p" else None
            if properties is not None and properties.find(f"{_WORD}sectPr") is not None:
                # A paragraph holding a section break keeps the section's page setup and stories.
                for part in list(child):
                    if part is not properties:
                        child.remove(part)
                continue
            self.body.remove(child)
        for relationship in self.document.part.rels.values():
            if relationship.reltype == _COMMENTS_REL and not relationship.is_external:
                comments = relationship.target_part.element
                for comment in list(comments):
                    comments.remove(comment)

    def replace_text(self, operation: dict[str, Any], log: EditLog) -> None:
        from docx.oxml.ns import qn

        count = 0
        for element in list(self.body.iter(qn("w:p"))):
            found = sum(
                replace_in_runs(
                    runs,
                    operation["find"],
                    operation["replace"],
                    match_case=operation["match_case"],
                )
                for runs in self._run_segments(element)
            )
            if found:
                count += found
                self._touch(element)
        if not count:
            raise OperationError("The text to find wasn't found in the body or its tables.")
        log.change(f"Replaced {count} matches.", count=count)

    def insert_paragraphs(self, operation: dict[str, Any], log: EditLog) -> None:
        elements = []
        for spec in operation["paragraphs"]:
            spec = {"runs": [spec]} if isinstance(spec, str) else spec
            paragraph = self._new_paragraph()
            if spec.get("style"):
                paragraph.style = self._style(spec["style"], "paragraph")
            _write_runs(paragraph, spec["runs"], base=None)
            elements.append(paragraph._p)
        self._insert(elements, operation)
        log.change(f"Inserted {len(elements)} paragraphs.")

    def set_paragraph(self, operation: dict[str, Any], log: EditLog) -> None:
        paragraph = self._paragraph_at(operation)
        if operation.get("runs") is not None:
            if not _simple_paragraph(paragraph._p):
                raise OperationError(
                    f"Paragraph {operation['index']} has images, fields, links, or comments that "
                    "set_paragraph would remove. Use replace_text instead."
                )
            _rewrite_runs(paragraph, operation["runs"])
        if operation.get("style"):
            paragraph.style = self._style(operation["style"], "paragraph")
        self._touch(paragraph._p)
        log.change(f"Set paragraph {operation['index']}.")

    def delete_paragraphs(self, operation: dict[str, Any], log: EditLog) -> None:
        self._paragraph_at(operation)
        index, count = operation["index"], operation["count"]
        chosen = self._paragraphs()[index : index + count]
        if len(chosen) < count:
            raise OperationError(f"Only {len(chosen)} paragraphs start at index {index}.")
        for element in chosen:
            if element.pPr is not None and element.pPr.sectPr is not None:
                raise OperationError(
                    "A paragraph to delete ends a section, so deleting it would merge sections."
                )
        comments = self._ranges_within(chosen)
        for element in chosen:
            self.body.remove(element)
            self.touched.pop(element, None)
        self._remove_comments(comments)
        log.change(f"Deleted {count} paragraphs from index {index}.")

    def insert_table(self, operation: dict[str, Any], log: EditLog) -> None:
        rows = operation["rows"]
        columns = max(len(row) for row in rows)
        if not columns:
            raise OperationError("The table needs at least one column.")
        style = operation.get("style")
        table_style = self._style(style, "table") if style else self._default_table_style()
        table = self.document.add_table(rows=len(rows), cols=columns)
        if table_style is not None:
            table.style = table_style
        for row_index, values in enumerate(rows):
            for column, text in enumerate(values):
                table.cell(row_index, column).text = text
        self.body.remove(table._tbl)
        self._insert([table._tbl], operation)
        log.change(f"Inserted a {len(rows)}-row table.")

    def set_table_cells(self, operation: dict[str, Any], log: EditLog) -> None:
        from docx.oxml.ns import qn
        from docx.table import Table

        tables = list(self.body.iterchildren(qn("w:tbl")))
        index = operation["table_index"]
        if not 0 <= index < len(tables):
            raise OperationError(
                f"The document has {len(tables)} body tables; index {index} doesn't exist."
            )
        element = tables[index]
        table = Table(element, self.document.part)
        start_row, start_column = operation["start_row"], operation["start_column"]
        values = operation["values"]
        while len(element.tr_lst) < start_row + len(values):
            element.tr_lst[-1].addnext(_blank_row(element.tr_lst[-1]))
        writes = []
        for row_offset, row_values in enumerate(values):
            row = table.rows[start_row + row_offset]
            cells = row.cells
            if start_column + len(row_values) > len(cells):
                raise OperationError(
                    f"Row {start_row + row_offset} of table {index} has {len(cells)} cells."
                )
            for column_offset, text in enumerate(row_values):
                position = start_column + column_offset
                cell = cells[position]
                # A merged cell repeats across the positions it covers.
                covered = (position and cells[position - 1]._tc is cell._tc) or (
                    cell._tc.getparent() is not row._tr
                )
                if covered and text:
                    raise OperationError(
                        f"Row {start_row + row_offset} of table {index} has text for column "
                        f"{position}, which a merged cell covers. Write to the merged cell's "
                        "first position only."
                    )
                if not covered:
                    writes.append((cell, text))
        for cell, text in writes:
            _set_cell_text(cell, text)
        self._touch(element)
        log.change(f"Set {sum(len(row) for row in values)} cells of table {index}.")

    def add_image(self, operation: dict[str, Any], log: EditLog) -> None:
        from docx.shared import Pt

        index = operation["image"]
        if not 0 <= index < len(self.images):
            raise OperationError("The image File wasn't loaded.")
        paragraph = self._new_paragraph()
        self._insert([paragraph._p], operation)
        width = Pt(operation["width"]) if operation.get("width") is not None else None
        try:
            picture = paragraph.add_run().add_picture(io.BytesIO(self.images[index]), width=width)
        except Exception:
            raise OperationError("The image File isn't a readable PNG or JPEG image.") from None
        section = self._section_of(paragraph._p)
        room = section.page_width - section.left_margin - section.right_margin
        if room > 0 and picture.width > room:
            picture.height = int(picture.height * room / picture.width)
            picture.width = room
        log.change("Added an image.")

    def page_break(self, operation: dict[str, Any], log: EditLog) -> None:
        from docx.enum.text import WD_BREAK

        paragraph = self._new_paragraph()
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        self._insert([paragraph._p], operation)
        log.change("Added a page break.")

    def set_story(self, operation: dict[str, Any], log: EditLog) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph

        number, variant = operation["section"], operation["variant"]
        story = self._own_story(operation["op"], number, variant)
        container = story._element
        paragraphs = list(container.iterchildren(qn("w:p")))
        replaceable = [element for element in paragraphs if _simple_paragraph(element)]
        first = replaceable[0] if replaceable else None
        base_paragraph = (
            deepcopy(first.pPr) if first is not None and first.pPr is not None else None
        )
        base_run = _base_run_properties(first) if first is not None else None
        for line in operation["text"].split("\n"):
            element = OxmlElement("w:p")
            if base_paragraph is not None:
                element.append(deepcopy(base_paragraph))
            if line:
                _write_runs(Paragraph(element, story), [line], base=base_run)
            if first is not None:
                first.addprevious(element)
            else:
                container.append(element)
        for element in replaceable:
            container.remove(element)
        label = operation["op"].removeprefix("set_")
        kept = len(paragraphs) - len(replaceable) + len(list(container.iterchildren(qn("w:tbl"))))
        if kept:
            log.warn(
                f"The {label} of section {number} kept {kept} paragraphs or tables with images, "
                "fields, or page numbers."
            )
        self.stories[(operation["op"], number, variant)] = None
        log.change(f"Set the {variant.replace('_', ' ')} {label} of section {number}.")

    def _own_story(self, op: str, number: int, variant: str) -> Any:
        """Returns a section's own header or footer, unlinking it from the earlier section first."""
        sections = self.document.sections
        if not 1 <= number <= len(sections):
            raise OperationError(f"The document has {len(sections)} sections.")
        section = sections[number - 1]
        if variant == "first_page":
            section.different_first_page_header_footer = True
        if variant == "even_page" and not self.document.settings.odd_and_even_pages_header_footer:
            raise OperationError(
                "Even-page headers and footers are off in this document. Use the primary variant."
            )
        attribute = _STORIES[(op, variant)]
        story = getattr(section, attribute)
        if story.is_linked_to_previous:
            inherited = _inherited_story(sections[: number - 1], attribute)
            story.is_linked_to_previous = False
            if inherited is not None:
                # The section showed the earlier story, so it starts as a copy of it.
                _copy_story(inherited, story)
        return story

    def add_comment(self, operation: dict[str, Any], log: EditLog) -> None:
        paragraph = self._paragraph_at(operation)
        runs = paragraph.runs
        if not runs:
            raise OperationError(f"Paragraph {operation['index']} has no text to comment on.")
        comment = self.document.add_comment(
            (runs[0], runs[-1]),
            text=operation["text"],
            author=self.author,
            initials=_initials(self.author),
        )
        self._touch(paragraph._p)
        self.comment_ids.add(comment.comment_id)
        log.change(f"Commented on paragraph {operation['index']}.", comment_id=comment.comment_id)

    def end_with_paragraph(self) -> None:
        """Adds an empty paragraph when the body would otherwise end with a table or be empty."""
        blocks = [child for child in self.body if child.tag != f"{_WORD}sectPr"]
        if not blocks or blocks[-1].tag != f"{_WORD}p":
            self._append(self._new_paragraph()._p)

    def readback(self, log: EditLog) -> dict[str, Any]:
        groups: dict[str, Iterable[dict[str, Any]]] = {
            "blocks": self._touched_blocks(log),
            "headers": self._touched_stories("set_header", log),
            "footers": self._touched_stories("set_footer", log),
            "comments": (
                read_comment(comment, log.page)
                for comment in self.document.comments
                if comment.comment_id in self.comment_ids
            ),
        }
        result: dict[str, Any] = {}
        for key, items in groups.items():
            kept = []
            for item in items:
                if not log.page.fits(item):
                    result[f"{key}_truncated"] = True
                    break
                kept.append(item)
            if kept or key == "blocks":
                result[key] = kept
        return result

    def _section_of(self, element: Any) -> Any:
        """Returns the section a body paragraph belongs to: the next section break's."""
        from docx.section import Section

        for block in (element, *element.itersiblings()):
            properties = block.find(f"{_WORD}pPr") if block.tag == f"{_WORD}p" else None
            if properties is not None and properties.find(f"{_WORD}sectPr") is not None:
                return Section(properties.find(f"{_WORD}sectPr"), self.document.part)
        return self.document.sections[-1]

    def _ranges_within(self, chosen: list[Any]) -> set[str]:
        """Returns comment ids wholly inside the chosen blocks, failing when a range crosses out.

        A bookmark or comment with one end inside and one outside would be left
        pointing at the wrong text.
        """
        inside: dict[tuple[str, str], None] = {}
        for element in chosen:
            for node in element.iter(*_RANGE_ENDS):
                inside[(_RANGE_ENDS[node.tag], node.get(f"{_WORD}id"))] = None
        if not inside:
            return set()
        kept = {id(element) for element in chosen}
        for block in self.body:
            if id(block) in kept:
                continue
            for node in block.iter(*_RANGE_ENDS):
                kind = _RANGE_ENDS[node.tag]
                if (kind, node.get(f"{_WORD}id")) in inside:
                    raise OperationError(
                        f"A {kind} covers text inside and outside the paragraphs to delete, so "
                        "deleting them would break it. Delete the whole range, or use "
                        "replace_text instead."
                    )
        return {range_id for kind, range_id in inside if kind == "comment"}

    def _remove_comments(self, comment_ids: set[str]) -> None:
        """Removes comments whose whole range was deleted, so none is left without its text."""
        if not comment_ids:
            return
        for relationship in self.document.part.rels.values():
            if relationship.reltype == _COMMENTS_REL and not relationship.is_external:
                comments = relationship.target_part.element
                for comment in list(comments):
                    if comment.get(f"{_WORD}id") in comment_ids:
                        comments.remove(comment)

    def _touched_stories(self, op: str, log: EditLog) -> Iterator[dict[str, Any]]:
        sections = self.document.sections
        for story_op, number, variant in self.stories:
            if story_op == op:
                story = getattr(sections[number - 1], _STORIES[(op, variant)])
                yield read_story(story, number, variant, log.page)

    def _touched_blocks(self, log: EditLog) -> Iterator[dict[str, Any]]:
        """Yields touched body blocks in document order, numbered as reads number them."""
        counts = {"p": 0, "tbl": 0}
        for child in self.body:
            if child.tag not in _BLOCKS:
                continue
            tag = child.tag.removeprefix(_WORD)
            index = counts.get(tag, 0)
            if tag in counts:
                counts[tag] += 1
            if child in self.touched:
                yield read_block(self.document, child, tag, index, log.page)

    def _paragraphs(self) -> list[Any]:
        from docx.oxml.ns import qn

        return list(self.body.iterchildren(qn("w:p")))

    def _paragraph_at(self, operation: dict[str, Any]) -> Any:
        from docx.text.paragraph import Paragraph

        paragraphs = self._paragraphs()
        index, expected = operation["index"], operation["expect_text"]
        if not 0 <= index < len(paragraphs):
            raise OperationError(
                f"The document has {len(paragraphs)} body paragraphs; index {index} doesn't exist."
            )
        paragraph = Paragraph(paragraphs[index], self.document.part)
        text = paragraph.text
        if not text.startswith(expected) or (not expected and text):
            # The current text is file content, so it stays out of the message.
            raise OperationError(
                f"Paragraph {index} doesn't start with expect_text. Read the document again "
                "for current indexes."
            )
        return paragraph

    def _new_paragraph(self) -> Any:
        from docx.oxml import OxmlElement
        from docx.text.paragraph import Paragraph

        return Paragraph(OxmlElement("w:p"), self.document.part)

    def _append(self, element: Any) -> None:
        end = self.body.sectPr
        if end is not None:
            end.addprevious(element)
        else:
            self.body.append(element)

    def _insert(self, elements: list[Any], operation: dict[str, Any]) -> None:
        if operation.get("index") is None:
            for element in elements:
                self._append(element)
        else:
            anchor = self._paragraph_at(operation)._p
            for element in elements if operation["position"] == "before" else reversed(elements):
                if operation["position"] == "before":
                    anchor.addprevious(element)
                else:
                    anchor.addnext(element)
        for element in elements:
            self._touch(element)

    def _run_segments(self, element: Any) -> Iterator[list[Any]]:
        """Yields runs whose text can change safely, split at links, fields, and other content."""
        from docx.text.paragraph import Paragraph
        from docx.text.run import Run

        paragraph = Paragraph(element, self.document.part)
        segment: list[Any] = []
        for child in element:
            if child.tag == f"{_WORD}r" and _replaceable_run(child):
                segment.append(Run(child, paragraph))
                continue
            if child.tag in _MARKERS or _comment_reference(child):
                continue
            if segment:
                yield segment
                segment = []
            if child.tag == f"{_WORD}hyperlink":
                links = [Run(run, paragraph) for run in child if run.tag == f"{_WORD}r"]
                if links and all(_replaceable_run(run._r) for run in links):
                    yield links
        if segment:
            yield segment

    def _style(self, style_name: str, kind: str) -> Any:
        from docx.enum.style import WD_STYLE_TYPE

        style_type = WD_STYLE_TYPE.PARAGRAPH if kind == "paragraph" else WD_STYLE_TYPE.TABLE
        styles = [style for style in self.document.styles if style.type == style_type]
        for style in styles:
            if style.name == style_name:
                return style
        for style in styles:
            if (style.name or "").casefold() == style_name.casefold():
                return style
        names = ", ".join(repr((style.name or "")[:100]) for style in styles[:40])
        raise OperationError(
            f"The document has no {kind} style {style_name!r}. Its styles are: {names}."
        )

    def _default_table_style(self) -> Any:
        from docx.enum.style import WD_STYLE_TYPE

        for style in self.document.styles:
            if style.type == WD_STYLE_TYPE.TABLE and style.name == "Table Grid":
                return style
        return None

    def _touch(self, element: Any) -> None:
        while element is not None and element.getparent() is not self.body:
            element = element.getparent()
        if element is not None:
            self.touched[element] = None


def _write_runs(paragraph: Any, runs: list[Any], *, base: Any) -> None:
    for spec in runs:
        spec = {"text": spec} if isinstance(spec, str) else spec
        run = paragraph.add_run()
        if base is not None:
            run._r.insert(0, deepcopy(base))
        run.text = spec["text"]
        _format_run(run, spec)


def _format_run(run: Any, spec: dict[str, Any]) -> None:
    from docx.shared import Pt

    for key in ("bold", "italic", "underline"):
        if spec.get(key) is not None:
            setattr(run, key, spec[key])
    if spec.get("size") is not None:
        run.font.size = Pt(spec["size"])
    if spec.get("color"):
        _set_color(run.font.color, spec["color"])


def _set_color(color_format: Any, value: str) -> None:
    from docx.enum.dml import MSO_THEME_COLOR
    from docx.shared import RGBColor

    if value.startswith("#"):
        color_format.rgb = RGBColor.from_string(value[1:].upper())
        return
    try:
        color_format.theme_color = MSO_THEME_COLOR[value.split(":", 1)[1].upper()]
    except KeyError:
        raise OperationError(f"{value!r} isn't a theme colour.") from None


def _base_run_properties(paragraph: Any) -> Any:
    """Returns the first run's formatting, without tracked-change history."""
    for child in paragraph:
        if child.tag == f"{_WORD}r":
            properties = child.find(f"{_WORD}rPr")
            if properties is None:
                return None
            copy = deepcopy(properties)
            for change in copy.findall(f"{_WORD}rPrChange"):
                copy.remove(change)
            return copy
    return None


def _simple_run(run: Any) -> bool:
    """Checks whether a run holds only text, tabs, and line breaks, so rewriting it loses nothing."""
    for child in run:
        if child.tag == f"{_WORD}br":
            if child.get(f"{_WORD}type") not in (None, "textWrapping"):
                return False
        elif child.tag not in _SIMPLE_RUN_PARTS:
            return False
    return True


def _replaceable_run(run: Any) -> bool:
    return _simple_run(run) and not any(child.tag in _SEMANTIC_RUN_PARTS for child in run)


def _inherited_story(earlier_sections: Any, attribute: str) -> Any:
    """Returns the nearest earlier section's own header or footer of a kind, or None."""
    for section in reversed(earlier_sections):
        story = getattr(section, attribute)
        if not story.is_linked_to_previous:
            return story
    return None


def _copy_story(source: Any, target: Any) -> None:
    """Replaces a header or footer's content with a copy of another's, sharing its images."""
    container = target._element
    for child in list(container):
        container.remove(child)
    copied = [deepcopy(child) for child in source._element]
    for child in copied:
        container.append(child)
    mapping: dict[str, str] = {}
    for element in copied:
        for node in element.iter():
            for key, value in node.attrib.items():
                if not key.startswith(_R_NAMESPACE):
                    continue
                if value not in mapping:
                    mapping[value] = _relate_copy(source.part, target.part, value)
                node.set(key, mapping[value])


def _relate_copy(source: Any, target: Any, rel_id: str) -> str:
    relationship = source.rels.get(rel_id)
    if relationship is not None and relationship.is_external:
        return target.relate_to(relationship.target_ref, relationship.reltype, is_external=True)
    if relationship is not None and relationship.reltype == _IMAGE_REL:
        return target.relate_to(relationship.target_part, relationship.reltype)
    raise OperationError(
        "The header or footer this section shows from an earlier section has an embedded "
        "object that can't be copied. Set it in the earlier section instead."
    )


def _comment_reference(run: Any) -> bool:
    return run.tag == f"{_WORD}r" and {child.tag for child in run} <= {
        f"{_WORD}rPr",
        f"{_WORD}commentReference",
    }


def _simple_paragraph(paragraph: Any) -> bool:
    return all(
        child.tag in _SIMPLE_PARAGRAPH_PARTS or (child.tag == f"{_WORD}r" and _simple_run(child))
        for child in paragraph
    ) and (paragraph.pPr is None or paragraph.pPr.sectPr is None)


def _set_cell_text(cell: Any, text: str) -> None:
    from docx.text.paragraph import Paragraph

    element = cell._tc
    paragraphs = element.p_lst
    if element.tbl_lst or not all(_simple_paragraph(paragraph) for paragraph in paragraphs):
        raise OperationError(
            "A cell has images, fields, links, or a nested table that setting its text would remove."
        )
    first = paragraphs[0]
    if any(next(extra.iter(*_BOOKMARK_ENDS), None) is not None for extra in paragraphs[1:]):
        raise OperationError("A cell has a bookmark that setting its text would remove.")
    for extra in paragraphs[1:]:
        element.remove(extra)
    _rewrite_runs(Paragraph(first, cell), [text] if text else [])


def _rewrite_runs(paragraph: Any, runs: list[Any]) -> None:
    """Replaces a simple paragraph's text with new runs in the place of the old ones.

    New runs take the first old run's formatting. A bookmark around all of the
    text keeps covering it; one around only part of it fails the operation.
    """
    element = paragraph._p
    _check_partial_bookmarks(element)
    base = _base_run_properties(element)
    children = list(element)
    first_run = next(
        (index for index, child in enumerate(children) if child.tag == f"{_WORD}r"), None
    )
    removed = (f"{_WORD}r", f"{_WORD}proofErr")
    anchor = None
    if first_run is not None:
        anchor = next(
            (child for child in reversed(children[:first_run]) if child.tag not in removed), None
        )
    for child in children:
        if child.tag in removed:
            element.remove(child)
    before = set(map(id, element))
    _write_runs(paragraph, runs, base=base)
    if first_run is None:
        return
    for run in [child for child in element if id(child) not in before]:
        if anchor is None:
            element.insert(0, run)
        else:
            anchor.addnext(run)
        anchor = run


def _check_partial_bookmarks(paragraph: Any) -> None:
    """Fails when a bookmark end sits between runs, unless the bookmark covers no text."""
    children = list(paragraph)
    runs = [index for index, child in enumerate(children) if child.tag == f"{_WORD}r"]
    if len(runs) < 2:
        return
    positions: dict[str, list[int]] = {}
    for index, child in enumerate(children):
        if child.tag in _BOOKMARK_ENDS:
            positions.setdefault(child.get(f"{_WORD}id"), []).append(index)
    for ends in positions.values():
        empty = len(ends) == 2 and not any(ends[0] < run < ends[1] for run in runs)
        if not empty and any(runs[0] < end < runs[-1] for end in ends):
            raise OperationError(
                "A bookmark covers part of this text, and rewriting it would move the bookmark. "
                "Use replace_text instead."
            )


def _blank_row(row: Any) -> Any:
    """Returns a copy of a table row's row, cell, paragraph, and run formatting, without content.

    Bookmarks, comments, images, and fields stay in the original row, and a
    vertical merge ends at it.
    """
    from docx.oxml import OxmlElement

    blank = OxmlElement("w:tr")
    for child in row:
        if child.tag in (f"{_WORD}tblPrEx", f"{_WORD}trPr"):
            properties = deepcopy(child)
            for part in list(properties):
                if part.tag in _UNCOPIED_ROW_PARTS:
                    properties.remove(part)
            blank.append(properties)
        elif child.tag == f"{_WORD}tc":
            blank.append(_blank_cell(child))
        else:
            raise OperationError(
                "The table's last row has content controls or tracked changes, so new rows can't "
                "copy its formatting."
            )
    return blank


def _blank_cell(cell: Any) -> Any:
    from docx.oxml import OxmlElement

    blank = OxmlElement("w:tc")
    properties = cell.find(f"{_WORD}tcPr")
    if properties is not None:
        properties = deepcopy(properties)
        for merge in properties.findall(f"{_WORD}vMerge"):
            properties.remove(merge)
        blank.append(properties)
    paragraph = OxmlElement("w:p")
    source = cell.find(f"{_WORD}p")
    if source is not None:
        if source.find(f"{_WORD}pPr") is not None:
            paragraph.append(deepcopy(source.find(f"{_WORD}pPr")))
        run_properties = _base_run_properties(source)
        if run_properties is not None:
            run = OxmlElement("w:r")
            run.append(run_properties)
            paragraph.append(run)
    blank.append(paragraph)
    return blank


def _initials(author: str) -> str:
    return "".join(word[0] for word in author.split() if word)[:4].upper()
