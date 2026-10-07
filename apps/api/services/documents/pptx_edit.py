# apps/api/services/documents/pptx_edit.py

"""Edits to presentations. Runs only inside the document worker."""

import io
from collections.abc import Iterator
from copy import deepcopy
from typing import Any

from services.documents.editing import (
    R_NAMESPACE,
    EditLog,
    OperationError,
    format_font,
    relink,
    replace_in_runs,
    run_operations,
)
from services.documents.pptx_model import read_slide
from services.documents.reading import name

_EMU_PER_POINT = 12_700
_GROWTH_RATIO = 1.5
_GROWTH_MIN_CHARS = 40
_ASPECT_TOLERANCE = 0.05
_CHART_TYPES = {
    "column": "COLUMN_CLUSTERED",
    "stacked_column": "COLUMN_STACKED",
    "bar": "BAR_CLUSTERED",
    "stacked_bar": "BAR_STACKED",
    "line": "LINE",
    "line_markers": "LINE_MARKERS",
    "pie": "PIE",
    "doughnut": "DOUGHNUT",
    "area": "AREA",
}
_SINGLE_SERIES = frozenset({"pie", "doughnut"})
_LINKS = ("hlinkClick", "hlinkMouseOver")
_SECTION_SLIDE = "{http://schemas.microsoft.com/office/powerpoint/2010/main}sldId"
# PowerPoint's per-slide identity, which a copy mustn't share with its source.
_CREATION_ID_EXTENSION = "{BB962C8B-B14F-4D97-AF65-F5344CB8AC3E}"


def edit_presentation(
    presentation: Any, args: dict[str, Any], images: list[bytes]
) -> dict[str, Any]:
    """Applies operations to an open deck and returns changes, warnings, and read-back.

    With `clear`, the deck's slides are removed first so only its layouts remain.
    """
    log = EditLog(args)
    editor = _PresentationEditor(presentation, images)
    if args.get("clear"):
        for slide in list(presentation.slides):
            editor.remove_slide(slide.slide_id)
    run_operations(args["operations"], editor.handlers(), log)
    return log.result(lambda: editor.readback(log))


class _PresentationEditor:
    def __init__(self, presentation: Any, images: list[bytes]) -> None:
        self.presentation = presentation
        self.images = images
        # Slide ids in the order operations touched them, for the read-back.
        self.touched: dict[int, None] = {}

    def handlers(self) -> dict[str, Any]:
        return {
            "add_slide": self.add_slide,
            "duplicate_slide": self.duplicate_slide,
            "delete_slide": self.delete_slide,
            "move_slide": self.move_slide,
            "set_text": self.set_text,
            "replace_text": self.replace_text,
            "add_text_box": self.add_text_box,
            "set_table": self.set_table,
            "add_table": self.add_table,
            "add_image": self.add_image,
            "replace_image": self.replace_image,
            "add_chart": self.add_chart,
            "set_chart_data": self.set_chart_data,
            "set_notes": self.set_notes,
            "set_geometry": self.set_geometry,
            "delete_shape": self.delete_shape,
        }

    def add_slide(self, operation: dict[str, Any], log: EditLog) -> None:
        layout = self._layout(operation["layout"])
        available = {shape.placeholder_format.idx for shape in layout.placeholders}
        for placeholder in operation.get("placeholders") or []:
            if placeholder["idx"] not in available:
                raise OperationError(
                    f"Layout {operation['layout']!r} has no placeholder idx {placeholder['idx']}. "
                    f"Its placeholders are: {sorted(available)}."
                )
        slide = self.presentation.slides.add_slide(layout)
        if operation.get("position") is not None:
            self._place(slide, operation["position"])
        for placeholder in operation.get("placeholders") or []:
            shape = slide.placeholders[placeholder["idx"]]
            if not shape.has_text_frame:
                raise OperationError(f"Placeholder idx {placeholder['idx']} doesn't hold text.")
            _write_paragraphs(shape.text_frame, placeholder["paragraphs"])
        if operation.get("notes"):
            self._set_notes(slide, operation["notes"])
        self._touch(slide)
        log.change(
            f"Added slide {self._number(slide)} from layout {name(layout.name)!r}.",
            slide_id=slide.slide_id,
        )

    def duplicate_slide(self, operation: dict[str, Any], log: EditLog) -> None:
        source = self._slide(operation["slide_id"])
        copy = self.presentation.slides.add_slide(source.slide_layout)
        tree = copy.shapes._spTree
        for child in _shape_elements(tree):
            tree.remove(child)
        copied = [deepcopy(child) for child in _shape_elements(source.shapes._spTree)]
        for child in copied:
            tree.append(child)
        background = source._element.cSld.bg
        if background is not None:
            copied.append(deepcopy(background))
            copy._element.cSld.insert(0, copied[-1])
        copied.extend(_copy_slide_settings(source._element, copy._element))
        _relink(copied, source.part, copy.part)
        if source.has_notes_slide:
            _copy_notes(source.notes_slide, copy.notes_slide)
        position = operation.get("position") or self._number(source) + 1
        self._place(copy, position)
        self._touch(copy)
        log.change(
            f"Copied slide {self._number(source)} to slide {self._number(copy)}.",
            slide_id=copy.slide_id,
        )

    def delete_slide(self, operation: dict[str, Any], log: EditLog) -> None:
        slide = self._slide(operation["slide_id"])
        number = self._number(slide)
        self.remove_slide(slide.slide_id)
        log.change(f"Deleted slide {number}.", slide_id=operation["slide_id"])

    def remove_slide(self, slide_id: int) -> None:
        slide_ids = self.presentation.slides._sldIdLst
        element = next(item for item in slide_ids.sldId_lst if item.id == slide_id)
        slide_ids.remove(element)
        # PowerPoint repairs a deck whose sections still list a removed slide.
        for entry in list(self.presentation._element.iter(_SECTION_SLIDE)):
            if entry.get("id") == str(slide_id):
                entry.getparent().remove(entry)
        _drop_unused(self.presentation.part, element.rId)
        self.touched.pop(slide_id, None)

    def move_slide(self, operation: dict[str, Any], log: EditLog) -> None:
        slide = self._slide(operation["slide_id"])
        self._place(slide, operation["position"])
        self._touch(slide)
        log.change(f"Moved the slide to position {operation['position']}.", slide_id=slide.slide_id)

    def set_text(self, operation: dict[str, Any], log: EditLog) -> None:
        slide, shape = self._shape(operation)
        if not shape.has_text_frame:
            raise OperationError(
                f"Shape {shape.shape_id} has no text. Use set_table for tables and "
                "set_chart_data for charts."
            )
        before = len(shape.text_frame.text)
        _write_paragraphs(shape.text_frame, operation["paragraphs"])
        after = len(shape.text_frame.text)
        if before and after > before * _GROWTH_RATIO and after - before >= _GROWTH_MIN_CHARS:
            log.warn(
                f"Shape {shape.shape_id} on slide {self._number(slide)} grew from {before} to "
                f"{after} characters. Check it still fits, or split the slide."
            )
        self._touch(slide)
        log.change(
            f"Set the text of shape {shape.shape_id} on slide {self._number(slide)}.",
            slide_id=slide.slide_id,
            shape_id=shape.shape_id,
        )

    def replace_text(self, operation: dict[str, Any], log: EditLog) -> None:
        if operation.get("slide_id") is None:
            slides = list(self.presentation.slides)
        else:
            slides = [self._slide(operation["slide_id"])]
        count = 0
        for slide in slides:
            if operation.get("shape_id") is None:
                shapes = list(_all_shapes(slide.shapes))
            else:
                shapes = [self._shape(operation)[1]]
            found = sum(
                replace_in_runs(
                    runs,
                    operation["find"],
                    operation["replace"],
                    match_case=operation["match_case"],
                )
                for shape in shapes
                for runs in _run_segments(shape)
            )
            if found:
                self._touch(slide)
                count += found
        if not count:
            raise OperationError("The text to find wasn't found.")
        log.change(f"Replaced {count} matches.", count=count)

    def add_text_box(self, operation: dict[str, Any], log: EditLog) -> None:
        slide = self._slide(operation["slide_id"])
        box = slide.shapes.add_textbox(*_box(operation))
        box.text_frame.word_wrap = True
        _write_paragraphs(box.text_frame, operation["paragraphs"])
        self._touch(slide)
        log.change(
            f"Added a text box to slide {self._number(slide)}.",
            slide_id=slide.slide_id,
            shape_id=box.shape_id,
        )

    def set_table(self, operation: dict[str, Any], log: EditLog) -> None:
        slide, shape = self._shape(operation)
        if not getattr(shape, "has_table", False):
            raise OperationError(f"Shape {shape.shape_id} isn't a table.")
        table = shape.table
        columns = len(table.columns)
        rows = operation["rows"]
        if any(len(row) > columns for row in rows):
            raise OperationError(f"The table has {columns} columns; a row has more cells.")
        rows_xml = table._tbl
        _check_row_spans(rows_xml.tr_lst, len(rows))
        while len(rows_xml.tr_lst) < len(rows):
            template = rows_xml.tr_lst[-1]
            template.addnext(deepcopy(template))
        for extra in rows_xml.tr_lst[len(rows) :]:
            rows_xml.remove(extra)
        for row_index, values in enumerate(rows):
            for column in range(columns):
                _set_cell_text(table.cell(row_index, column), _at(values, column), row_index)
        shape.height = sum(row.height for row in table.rows)
        self._touch(slide)
        log.change(
            f"Set a {len(rows)}-row table on slide {self._number(slide)}.",
            slide_id=slide.slide_id,
            shape_id=shape.shape_id,
        )

    def add_table(self, operation: dict[str, Any], log: EditLog) -> None:
        from pptx.oxml.ns import qn

        slide = self._slide(operation["slide_id"])
        rows = operation["rows"]
        columns = max(len(row) for row in rows)
        if not columns:
            raise OperationError("The table needs at least one column.")
        frame = slide.shapes.add_table(len(rows), columns, *_box(operation))
        style_id = _default_table_style(self.presentation)
        if style_id:
            frame.table._tbl.tblPr.find(qn("a:tableStyleId")).text = style_id
        for row_index, values in enumerate(rows):
            for column in range(columns):
                _set_cell_text(frame.table.cell(row_index, column), _at(values, column), row_index)
        self._touch(slide)
        log.change(
            f"Added a table to slide {self._number(slide)}.",
            slide_id=slide.slide_id,
            shape_id=frame.shape_id,
        )

    def add_image(self, operation: dict[str, Any], log: EditLog) -> None:
        slide = self._slide(operation["slide_id"])
        left, top = _emu(operation["left"]), _emu(operation["top"])
        width = _emu(operation["width"]) if operation.get("width") is not None else None
        height = _emu(operation["height"]) if operation.get("height") is not None else None
        try:
            picture = slide.shapes.add_picture(self._image(operation), left, top, width, height)
        except (OperationError, MemoryError):
            raise
        except Exception:
            raise OperationError("The image File isn't a readable PNG or JPEG image.") from None
        if width is None and height is None:
            _fit(picture, self.presentation, left, top)
        self._touch(slide)
        log.change(
            f"Added an image to slide {self._number(slide)}.",
            slide_id=slide.slide_id,
            shape_id=picture.shape_id,
        )

    def replace_image(self, operation: dict[str, Any], log: EditLog) -> None:
        from pptx.oxml.ns import qn

        slide, shape = self._shape(operation)
        blips = shape._element.xpath("./p:blipFill/a:blip")
        if not blips:
            raise OperationError(f"Shape {shape.shape_id} isn't a picture.")
        try:
            image_part, rel_id = slide.part.get_or_add_image_part(self._image(operation))
        except (OperationError, MemoryError):
            raise
        except Exception:
            raise OperationError("The image File isn't a readable PNG or JPEG image.") from None
        previous = blips[0].get(qn("r:embed"))
        blips[0].set(qn("r:embed"), rel_id)
        if previous and previous != rel_id:
            _drop_unused(slide.part, previous)
        # A crop sized for the old image would cut the new one.
        for crop in shape._element.xpath("./p:blipFill/a:srcRect"):
            crop.getparent().remove(crop)
        pixel_width, pixel_height = image_part._px_size
        if shape.width and shape.height and pixel_width and pixel_height:
            ratio = (pixel_width / pixel_height) / (shape.width / shape.height)
            if abs(ratio - 1) > _ASPECT_TOLERANCE:
                log.warn(
                    f"The new image on slide {self._number(slide)} has a different shape from "
                    "its frame, so it's stretched. Adjust it with set_geometry."
                )
        self._touch(slide)
        log.change(
            f"Replaced the image in shape {shape.shape_id} on slide {self._number(slide)}.",
            slide_id=slide.slide_id,
            shape_id=shape.shape_id,
        )

    def add_chart(self, operation: dict[str, Any], log: EditLog) -> None:
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

        slide = self._slide(operation["slide_id"])
        chart_type = operation["chart_type"]
        data = _chart_data(operation, single_series=chart_type in _SINGLE_SERIES)
        frame = slide.shapes.add_chart(
            XL_CHART_TYPE[_CHART_TYPES[chart_type]], *_box(operation), data
        )
        chart = frame.chart
        # Without a title, PowerPoint would show a one-series chart's series name as one.
        chart.has_title = bool(operation.get("title"))
        if chart.has_title:
            chart.chart_title.text_frame.text = operation["title"]
        chart.has_legend = len(operation["series"]) > 1 or chart_type in _SINGLE_SERIES
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        self._touch(slide)
        log.change(
            f"Added a {chart_type} chart to slide {self._number(slide)}.",
            slide_id=slide.slide_id,
            shape_id=frame.shape_id,
        )

    def set_chart_data(self, operation: dict[str, Any], log: EditLog) -> None:
        slide, shape = self._shape(operation)
        if not getattr(shape, "has_chart", False):
            raise OperationError(f"Shape {shape.shape_id} isn't a chart.")
        try:
            kind = shape.chart.chart_type.name
        except (NotImplementedError, ValueError, KeyError):
            kind = "OTHER"
        if kind == "OTHER" or kind.startswith(("XY_", "BUBBLE")):
            raise OperationError(
                f"Shape {shape.shape_id} isn't a category chart, so its data can't be replaced."
            )
        single = "PIE" in kind or "DOUGHNUT" in kind
        shape.chart.replace_data(_chart_data(operation, single_series=single))
        self._touch(slide)
        log.change(
            f"Replaced the data of chart {shape.shape_id} on slide {self._number(slide)}.",
            slide_id=slide.slide_id,
            shape_id=shape.shape_id,
        )

    def set_notes(self, operation: dict[str, Any], log: EditLog) -> None:
        slide = self._slide(operation["slide_id"])
        self._set_notes(slide, operation["text"])
        self._touch(slide)
        log.change(f"Set the notes of slide {self._number(slide)}.", slide_id=slide.slide_id)

    def set_geometry(self, operation: dict[str, Any], log: EditLog) -> None:
        slide, shape = self._shape(operation)
        fields = [
            key for key in ("left", "top", "width", "height") if operation.get(key) is not None
        ]
        if not fields:
            raise OperationError("Set left, top, width, or height.")
        for key in fields:
            setattr(shape, key, _emu(operation[key]))
        self._touch(slide)
        log.change(
            f"Moved or resized shape {shape.shape_id} on slide {self._number(slide)}.",
            slide_id=slide.slide_id,
            shape_id=shape.shape_id,
        )

    def delete_shape(self, operation: dict[str, Any], log: EditLog) -> None:
        slide, shape = self._shape(operation)
        element = shape._element
        rel_ids = [
            value
            for node in element.iter()
            for key, value in node.attrib.items()
            if key.startswith(R_NAMESPACE)
        ]
        element.getparent().remove(element)
        for rel_id in rel_ids:
            _drop_unused(slide.part, rel_id)
        self._touch(slide)
        log.change(
            f"Deleted shape {operation['shape_id']} from slide {self._number(slide)}.",
            slide_id=slide.slide_id,
            shape_id=operation["shape_id"],
        )

    def readback(self, log: EditLog) -> dict[str, Any]:
        slides = []
        for slide_id in self.touched:
            slide = self.presentation.slides.get(slide_id)
            item = read_slide(slide, self._number(slide), log.page)
            if not log.page.fits(item):
                return {"slides": slides, "slides_truncated": True}
            slides.append(item)
        return {"slides": slides}

    def _slide(self, slide_id: int) -> Any:
        slide = self.presentation.slides.get(slide_id)
        if slide is None:
            raise OperationError(
                f"The deck has no slide with slide_id {slide_id}. Read the deck for current ids."
            )
        return slide

    def _shape(self, operation: dict[str, Any]) -> tuple[Any, Any]:
        slide = self._slide(operation["slide_id"])
        shape_id = operation["shape_id"]
        for shape in _all_shapes(slide.shapes):
            if shape.shape_id == shape_id:
                return slide, shape
        raise OperationError(f"Slide {self._number(slide)} has no shape with shape_id {shape_id}.")

    def _layout(self, layout_name: str) -> Any:
        layouts = [
            layout for master in self.presentation.slide_masters for layout in master.slide_layouts
        ]
        for matches in (
            lambda layout: layout.name == layout_name,
            lambda layout: layout.name.casefold() == layout_name.casefold(),
        ):
            for layout in layouts:
                if matches(layout):
                    return layout
        names = ", ".join(repr(layout.name[:100]) for layout in layouts[:30])
        raise OperationError(f"The deck has no layout {layout_name!r}. Its layouts are: {names}.")

    def _image(self, operation: dict[str, Any]) -> io.BytesIO:
        index = operation["image"]
        if not 0 <= index < len(self.images):
            raise OperationError("The image File wasn't loaded.")
        return io.BytesIO(self.images[index])

    def _place(self, slide: Any, position: int) -> None:
        slide_ids = self.presentation.slides._sldIdLst
        elements = slide_ids.sldId_lst
        if not 1 <= position <= len(elements):
            raise OperationError(f"position must be between 1 and {len(elements)}.")
        element = next(item for item in elements if item.id == slide.slide_id)
        slide_ids.remove(element)
        slide_ids.insert(position - 1, element)

    def _number(self, slide: Any) -> int:
        return self.presentation.slides.index(slide) + 1

    def _set_notes(self, slide: Any, text: str) -> None:
        frame = slide.notes_slide.notes_text_frame
        if frame is None:
            raise OperationError("The deck's notes pages have no text area.")
        frame.text = text

    def _touch(self, slide: Any) -> None:
        self.touched[slide.slide_id] = None


def _write_paragraphs(text_frame: Any, paragraphs: list[Any]) -> None:
    """Replaces a text frame's paragraphs, carrying over the first paragraph's and run's formatting."""
    from pptx.text.text import _Paragraph

    body = text_frame._txBody
    existing = list(body.p_lst)
    base_paragraph = deepcopy(existing[0].pPr) if existing and existing[0].pPr is not None else None
    base_run = _base_run_properties(body)
    for element in existing:
        body.remove(element)
    for spec in paragraphs:
        spec = {"runs": [spec]} if isinstance(spec, str) else spec
        element = body.add_p()
        if base_paragraph is not None:
            element.insert(0, deepcopy(base_paragraph))
        paragraph = _Paragraph(element, text_frame)
        if spec.get("level") is not None:
            paragraph.level = spec["level"]
        for run_spec in spec["runs"]:
            run_spec = {"text": run_spec} if isinstance(run_spec, str) else run_spec
            for position, part in enumerate(run_spec["text"].split("\v")):
                if position:
                    paragraph.add_line_break()
                if part:
                    run = paragraph.add_run()
                    if base_run is not None:
                        run._r.insert(0, deepcopy(base_run))
                    run.text = part
                    _format_run(run.font, run_spec)


def _base_run_properties(body: Any) -> Any:
    """Returns the first run's properties, without links, to format replacement text."""
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn

    source = next(body.iter(qn("a:rPr")), None)
    if source is None:
        end = next(body.iter(qn("a:endParaRPr")), None)
        if end is None:
            return None
        source = parse_xml(f"<a:rPr {nsdecls('a')}/>")
        for key, value in end.attrib.items():
            source.set(key, value)
        for child in end:
            source.append(deepcopy(child))
    properties = deepcopy(source)
    for link in _LINKS:
        for element in properties.findall(qn(f"a:{link}")):
            properties.remove(element)
    return properties


def _format_run(font: Any, spec: dict[str, Any]) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.util import Pt

    format_font(font, spec, points=Pt, rgb=RGBColor, themes=MSO_THEME_COLOR)


def _copy_slide_settings(source: Any, target: Any) -> list[Any]:
    """Copies slide attributes, the colour map, transition, timing, and extensions to a copy.

    Returns the copied elements so their relationships can be relinked.
    """
    from pptx.oxml.ns import qn

    # Namespaced attributes, such as mc:Ignorable, depend on the source part's declarations.
    for key, value in source.attrib.items():
        if not key.startswith("{"):
            target.set(key, value)
    if source.cSld.get("name") is not None:
        target.cSld.set("name", source.cSld.get("name"))
    for child in list(target):
        if child.tag != qn("p:cSld"):
            target.remove(child)
    copied = []
    for child in source:
        if child.tag == qn("p:cSld"):
            continue
        item = deepcopy(child)
        for extension in item.findall(qn("p:ext")) if child.tag == qn("p:extLst") else []:
            if extension.get("uri") == _CREATION_ID_EXTENSION:
                item.remove(extension)
        target.append(item)
        copied.append(item)
    return copied


def _copy_notes(source: Any, target: Any) -> None:
    """Replaces a new notes page's shapes with copies of the source's, keeping their formatting."""
    tree = target.shapes._spTree
    for child in _shape_elements(tree):
        tree.remove(child)
    copied = [deepcopy(child) for child in _shape_elements(source.shapes._spTree)]
    for child in copied:
        tree.append(child)
    _relink(copied, source.part, target.part)


def _check_row_spans(rows: list[Any], count: int) -> None:
    """Refuses a row count change that would cut through or copy a vertically merged cell."""
    if count > len(rows) and any(cell.rowSpan > 1 or cell.vMerge for cell in rows[-1].tc_lst):
        raise OperationError(
            "The table's last row is part of a merged cell, so rows can't be added by copying it."
        )
    for index, row in enumerate(rows[:count]):
        if any(index + cell.rowSpan > count for cell in row.tc_lst):
            raise OperationError(
                f"Removing rows after row {count - 1} would cut through a merged cell. Unmerge it "
                "or keep those rows."
            )


def _set_cell_text(cell: Any, text: str, row: int) -> None:
    if getattr(cell, "is_spanned", False):
        if text:
            raise OperationError(f"Row {row} has text for a cell covered by a merged cell.")
        return
    _write_paragraphs(cell.text_frame, text.split("\n"))


def _run_segments(shape: Any) -> Iterator[list[Any]]:
    """Yields runs of each paragraph in a shape or its table cells, split at breaks and fields."""
    from pptx.oxml.ns import qn
    from pptx.text.text import _Run

    frames = [shape.text_frame] if shape.has_text_frame else []
    if getattr(shape, "has_table", False):
        frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
    for frame in frames:
        for paragraph in frame.paragraphs:
            segment: list[Any] = []
            for element in paragraph._element.content_children:
                if element.tag == qn("a:r"):
                    segment.append(_Run(element, paragraph))
                elif segment:
                    yield segment
                    segment = []
            if segment:
                yield segment


def _all_shapes(shapes: Any) -> Iterator[Any]:
    from pptx.shapes.group import GroupShape

    for shape in shapes:
        yield shape
        if isinstance(shape, GroupShape):
            yield from _all_shapes(shape.shapes)


def _shape_elements(tree: Any) -> list[Any]:
    from pptx.oxml.ns import qn

    group_properties = {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}
    return [child for child in tree if child.tag not in group_properties]


def _relink(elements: list[Any], source: Any, target: Any) -> None:
    relink(
        elements,
        source,
        target,
        refusal="The slide has a chart, video, or embedded object that can't be copied. "
        "Add a slide and rebuild it instead.",
    )


def _drop_unused(part: Any, rel_id: str) -> None:
    if part._rel_ref_count(rel_id) == 0:
        part.drop_rel(rel_id)


def _chart_data(operation: dict[str, Any], *, single_series: bool) -> Any:
    from pptx.chart.data import CategoryChartData

    categories = operation["categories"]
    series = operation["series"]
    if single_series and len(series) != 1:
        raise OperationError("Pie and doughnut charts take exactly one series.")
    data = CategoryChartData()
    data.categories = categories
    for item in series:
        if len(item["values"]) != len(categories):
            raise OperationError(
                f"Series {item['name']!r} has {len(item['values'])} values for "
                f"{len(categories)} categories."
            )
        data.add_series(item["name"], item["values"])
    return data


def _default_table_style(presentation: Any) -> str | None:
    """Returns the deck's default table style, so new tables match its template."""
    from pptx.opc.constants import RELATIONSHIP_TYPE
    from pptx.oxml import parse_xml

    try:
        part = presentation.part.part_related_by(RELATIONSHIP_TYPE.TABLE_STYLES)
    except KeyError:
        return None
    return parse_xml(part.blob).get("def")


def _fit(picture: Any, presentation: Any, left: int, top: int) -> None:
    """Scales a picture down, keeping its shape, so it stays on the slide."""
    room_width = presentation.slide_width - left
    room_height = presentation.slide_height - top
    if room_width <= 0 or room_height <= 0 or not picture.width or not picture.height:
        return
    scale = min(1.0, room_width / picture.width, room_height / picture.height)
    if scale < 1.0:
        picture.width = int(picture.width * scale)
        picture.height = int(picture.height * scale)


def _box(operation: dict[str, Any]) -> tuple[int, int, int, int]:
    return tuple(_emu(operation[key]) for key in ("left", "top", "width", "height"))  # type: ignore[return-value]


def _emu(points: float) -> int:
    return round(points * _EMU_PER_POINT)


def _at(values: list[str], index: int) -> str:
    return values[index] if index < len(values) else ""
